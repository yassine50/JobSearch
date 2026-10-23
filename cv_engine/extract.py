"""
Layout-aware CV text extraction.

Why this exists: plain `PyPDF2.extract_text()` interleaves the columns of designer CVs
(Canva / Europass / moderncv), glues headings onto neighbouring lines, leaks icon-font glyphs
("/envel⌢pe", "(cid:131)") and breaks LaTeX accents ("Ing´ enierie"). Everything downstream
depends on clean, ordered lines, so we rebuild the reading order from character geometry:

  1. pdfminer gives every text line with its bbox, font name and font size.
  2. Icon-font characters are dropped; ligatures / accents / letter-spacing are repaired.
  3. A vertical gutter search splits two-column pages; right-aligned dates in single-column
     CVs are recognised and kept on the same line (separated by a TAB).
  4. Fragments on the same baseline are merged left→right with "\t" between distant pieces,
     so the parser can tell "Role\tDate" apart from running prose.

Falls back to PyPDF2 when pdfminer is unavailable, and flags image-only (scanned) PDFs.
"""
from __future__ import annotations

import io
import re
import statistics
import unicodedata
from dataclasses import dataclass, field
from typing import Dict, List, Optional, Tuple

ICON_FONT_RE = re.compile(r"awesome|fontello|icomoon|material.?icons|glyphicon|entypo|ionicons|feather|linearicons|webdings", re.I)
SYMBOL_BULLET_FONT_RE = re.compile(r"^(?:[A-Z]{6}\+)?(?:symbol(?:mt)?|wingdings[-\w]*|zapfdingbats|dingbats)$", re.I)
BOLD_FONT_RE = re.compile(r"bold|black|heavy|semibold|demi|extrabold|-bd\b|,bd|bd$", re.I)

BULLET_GLYPHS = "•●▪■◦○◆◇❖➢➤►▸▹‣⁃∙·✓✔☐□❯→»➔➜"


@dataclass
class Line:
    text: str
    size: float = 10.0
    bold: bool = False
    x0: float = 0.0
    x1: float = 0.0
    top: float = 0.0          # distance from top of the document (pages stacked)
    page: int = 0
    col: int = 0              # -1 = spans columns, 0 = left/only column, 1 = right column
    links: List[str] = field(default_factory=list)
    caps: bool = False

    def __repr__(self) -> str:  # pragma: no cover - debug aid
        return f"Line(p{self.page} c{self.col} {self.size:.1f}{'B' if self.bold else ''} {self.text!r})"


@dataclass
class Document:
    lines: List[Line]
    body_size: float = 10.0
    links: List[str] = field(default_factory=list)
    meta: Dict = field(default_factory=dict)

    @property
    def text(self) -> str:
        return "\n".join(l.text.replace("\t", "  ") for l in self.lines)


# ─── Text normalisation ───────────────────────────────────────────────────────

_LIGATURES = {"ﬁ": "fi", "ﬂ": "fl", "ﬀ": "ff", "ﬃ": "ffi", "ﬄ": "ffl", "ﬅ": "ft", "ﬆ": "st"}
_COMBINING = {"´": "\u0301", "`": "\u0300", "ˆ": "\u0302", "^": "\u0302", "¨": "\u0308", "˜": "\u0303", "¸": "\u0327"}
_ACCENT_RE = re.compile(r"(\s?)([´`ˆ¨˜])\s?([A-Za-z])")


def _compose_accent(m) -> str:
    lead, acc, letter = m.group(1), m.group(2), m.group(3)
    keep_space = lead if acc in "´`" else ""
    return keep_space + unicodedata.normalize("NFC", letter + _COMBINING[acc])


_CEDILLA_RE = re.compile(r"([cC])\s?¸")
_FA_NAME_LEAK_RE = re.compile(
    r"((?<![\w:/.])/?envel[⌢o]pe|(?<![\w:/.])/?map♂?mar¶?er|♂¶ap-¶arker|(?<![\w:/.])♂?phone(?=\+?\d)|"
    r"(?<![\w:/.])/linkedin(?=[a-z@])|(?<![\w:/.])/github(?=[a-z@])|(?<![\w:/.])/(?:globe|home|phone|envelope|map)(?=\s|\d|$)|"
    r"Ὄ5|♂|¶|⌢)",
)

_KERNING_WORD_REPAIRS = [
    (re.compile(r"\bf\s+aster\b", re.I), "faster"),
    (re.compile(r"\bAls\s+o\b"), "Also"),
    (re.compile(r"\bals\s+o\b"), "also"),
    (re.compile(r"\ban\s+d\b"), "and"),
    (re.compile(r"\bGoogle\s+P\s+lay\b", re.I), "Google Play"),
    (re.compile(r"\bP\s+lay\s+Console\b", re.I), "Play Console"),
    (re.compile(r"\bAra\s+bic\b", re.I), "Arabic"),
    (re.compile(r"\bimpac\s+t\b", re.I), "impact"),
    (re.compile(r"\bdevel\s+oper(s?)\b", re.I), r"developer\1"),
    (re.compile(r"\barchitec\s+ture\b", re.I), "architecture"),
    (re.compile(r"\bperfor\s+mance\b", re.I), "performance"),
    (re.compile(r"\bapplica\s+tion(s?)\b", re.I), r"application\1"),
    (re.compile(r"\bintegra\s+tion\b", re.I), "integration"),
    (re.compile(r"\bprodu\s+ction\b", re.I), "production"),
    (re.compile(r"\bcolla\s+boration\b", re.I), "collaboration"),
    (re.compile(r"\brespo\s+nsiveness\b", re.I), "responsiveness"),
    (re.compile(r"\bmain\s+tained\b", re.I), "maintained"),
    (re.compile(r"\bstandar\s+dizing\b", re.I), "standardizing"),
    (re.compile(r"\bconfigu\s+rable\b", re.I), "configurable"),
    (re.compile(r"\bnotifica\s+tions\b", re.I), "notifications"),
    (re.compile(r"\bauthenti\s+cation\b", re.I), "authentication"),
    (re.compile(r"\bexp\s+erience\b", re.I), "experience"),
    (re.compile(r"\bman\s+agement\b", re.I), "management"),
    (re.compile(r"\benvir\s+onment\b", re.I), "environment"),
]


def _collapse_letterspacing(s: str) -> str:
    """'A Z E R   B E L A I D' -> 'AZER BELAID' (display fonts with tracking)."""
    stripped = s.strip()
    words = [w for w in re.split(r"\s{2,}", stripped) if w]
    if not words:
        return s
    if all(re.fullmatch(r"\S(?: \S)*", w) for w in words):
        singles = sum(len(w.split(" ")) for w in words)
        if singles >= 5 and singles >= 2 * len(words):
            return " ".join(w.replace(" ", "") for w in words)
    return s


def normalize_text(s: str) -> str:
    if not s:
        return ""
    for k, v in _LIGATURES.items():
        s = s.replace(k, v)
    s = re.sub(r"\(cid:\d+\)", "", s)
    s = re.sub(r"[\ue000-\uf8ff]", "", s)            # private-use (icon fonts)
    s = _FA_NAME_LEAK_RE.sub(" ", s)
    s = _CEDILLA_RE.sub(lambda m: unicodedata.normalize("NFC", m.group(1) + "\u0327"), s)
    s = _ACCENT_RE.sub(_compose_accent, s)
    s = unicodedata.normalize("NFC", s)
    s = s.replace("\u00a0", " ").replace("\u200b", "").replace("\u2009", " ").replace("\u202f", " ")
    s = s.replace("’", "'").replace("‘", "'").replace("“", '"').replace("”", '"')
    if "\t" not in s:
        s = _collapse_letterspacing(s)
    # Repair spaced hyphens inside compound words and URLs ("cross -platform" -> "cross-platform", "yassine -youssef" -> "yassine-youssef")
    s = re.sub(r"([a-zA-Z0-9])\s+-\s*([a-zA-Z0-9])", lambda m: f"{m.group(1)}-{m.group(2)}" if (m.group(2).islower() or "/" in s) else m.group(0), s)
    s = re.sub(r"([a-zA-Z0-9])-\s+([a-z]+)\b", r"\1-\2", s)
    for pat, repl in _KERNING_WORD_REPAIRS:
        s = pat.sub(repl, s)
    s = re.sub(r"\b(\d+\s+(?:years?|months?|ans?|mois)|and\s+\d+\s+(?:months?|years?))\s+\1\b", r"\1", s, flags=re.I)
    s = "\t".join(re.sub(r"[ ]{2,}", " ", p).strip() for p in s.split("\t"))
    # "Sousse , Tunisie" / "C , Swift ,Php" -> "Sousse, Tunisie" / "C, Swift, Php" (Canva/Word spacing)
    s = re.sub(r"(?<=\w) +([,;])", r"\1", s)
    s = re.sub(r"([,;])(?=[A-Za-zÀ-ÿ])", r"\1 ", s)
    return s.strip()


def _is_caps(text: str) -> bool:
    letters = [c for c in text if c.isalpha()]
    return len(letters) >= 3 and sum(c.isupper() for c in letters) / len(letters) > 0.85


# ─── PDF via pdfminer ─────────────────────────────────────────────────────────

def _pdf_links(pdf_bytes: bytes) -> List[Tuple[int, Tuple[float, float, float, float], str]]:
    out = []
    try:
        import PyPDF2
        reader = PyPDF2.PdfReader(io.BytesIO(pdf_bytes))
        for pi, page in enumerate(reader.pages):
            for a in (page.get("/Annots") or []):
                try:
                    o = a.get_object()
                    act = o.get("/A")
                    uri = act.get_object().get("/URI") if act else None
                    if not uri:
                        continue
                    rect = tuple(float(x) for x in o.get("/Rect", [0, 0, 0, 0]))
                    out.append((pi, rect, str(uri).strip()))
                except Exception:
                    continue
    except Exception:
        pass
    return out


def _fragments_from_pdfminer(pdf_bytes: bytes):
    from pdfminer.high_level import extract_pages
    from pdfminer.layout import LAParams, LTChar, LTTextContainer, LTTextLine, LTFigure, LTImage

    pages = []
    # all_texts=True: Canva / Figma / Illustrator exports wrap every text box in nested Form XObjects
    # (LTFigure). Without it, pdfminer leaves those chars un-analysed and the whole CV text is lost.
    params = LAParams(line_margin=0.3, char_margin=2.0, word_margin=0.18, boxes_flow=None, all_texts=True)
    for page in extract_pages(io.BytesIO(pdf_bytes), laparams=params):
        frags, images = [], []

        def walk(el):
            if isinstance(el, LTTextContainer):
                for ln in el:
                    if isinstance(ln, LTTextLine):
                        frags.append(_frag_from_line(ln, LTChar))
            elif isinstance(el, LTImage):
                images.append(el.bbox)
            elif isinstance(el, LTFigure):
                has_text = False
                for child in el:
                    if isinstance(child, (LTTextContainer, LTFigure)):
                        has_text = True
                    walk(child)
                if not has_text:
                    images.append(el.bbox)

        for el in page:
            walk(el)
        pages.append((page.bbox, [f for f in frags if f and f["text"].strip()], images))
    return pages


_OCR_CACHE: Dict[str, List] = {}


def _fragments_from_ocr(pdf_bytes: bytes):
    """Fallback layout-preserving OCR for image-only / scanned PDFs (macOS Vision + PDFKit)."""
    import copy
    import hashlib
    import json
    import os
    import subprocess
    import tempfile

    if not pdf_bytes or not os.path.exists("/usr/bin/swift"):
        return []
    digest = hashlib.sha256(pdf_bytes).hexdigest()
    if digest in _OCR_CACHE:
        return copy.deepcopy(_OCR_CACHE[digest])
    swift_code = """
import Foundation
import PDFKit
import Vision

let args = CommandLine.arguments
guard args.count > 1 else { exit(1) }
let url = URL(fileURLWithPath: args[1])
guard let doc = PDFDocument(url: url) else { exit(1) }
var outPages: [[String: Any]] = []
for pi in 0..<doc.pageCount {
    guard let page = doc.page(at: pi) else { continue }
    let b = page.bounds(for: .mediaBox)
    let w = b.width > 0 ? b.width : 595.0
    let h = b.height > 0 ? b.height : 842.0
    let img = page.thumbnail(of: CGSize(width: w * 2.0, height: h * 2.0), for: .mediaBox)
    guard let cg = img.cgImage(forProposedRect: nil, context: nil, hints: nil) else { continue }
    let req = VNRecognizeTextRequest()
    req.recognitionLevel = .accurate
    req.usesLanguageCorrection = true
    let handler = VNImageRequestHandler(cgImage: cg, options: [:])
    try? handler.perform([req])
    var frags: [[String: Any]] = []
    for obs in req.results ?? [] {
        guard let top = obs.topCandidates(1).first else { continue }
        let bb = obs.boundingBox
        let x0 = Double(bb.minX * w)
        let x1 = Double(bb.maxX * w)
        let y0 = Double(bb.minY * h)
        let y1 = Double(bb.maxY * h)
        let sz = max(8.0, min(22.0, (y1 - y0) * 0.78))
        frags.append(["text": top.string, "x0": x0, "x1": x1, "y0": y0, "y1": y1, "size": sz])
    }
    outPages.append(["w": Double(w), "h": Double(h), "frags": frags])
}
if let data = try? JSONSerialization.data(withJSONObject: outPages),
   let str = String(data: data, encoding: .utf8) {
    print(str)
}
"""
    tmp_path = None
    try:
        with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as tf:
            tf.write(pdf_bytes)
            tmp_path = tf.name
        proc = subprocess.run(
            ["/usr/bin/swift", "-e", swift_code, tmp_path],
            capture_output=True,
            text=True,
            timeout=25,
        )
        if proc.returncode != 0 or not proc.stdout.strip():
            return []
        raw_pages = json.loads(proc.stdout.strip())
        pages = []
        for p in raw_pages:
            w, h = float(p.get("w", 595.0)), float(p.get("h", 842.0))
            frags = []
            for f in p.get("frags", []):
                txt = (f.get("text") or "").strip()
                if not txt:
                    continue
                sz = float(f.get("size", 10.0))
                frags.append({
                    "text": txt,
                    "x0": float(f["x0"]),
                    "x1": float(f["x1"]),
                    "y0": float(f["y0"]),
                    "y1": float(f["y1"]),
                    "size": sz,
                    "bold": _is_caps(txt) or sz >= 11.8,
                })
            pages.append(((0.0, 0.0, w, h), frags, []))
        if sum(len(f["text"].strip()) for _, frs, _ in pages for f in frs) >= 80:
            _OCR_CACHE[digest] = copy.deepcopy(pages)
        return pages
    except Exception:
        return []
    finally:
        if tmp_path and os.path.exists(tmp_path):
            try:
                os.remove(tmp_path)
            except Exception:
                pass


def _frag_from_line(ln, LTChar):
    items = list(ln)
    chars = []
    real_seen = []
    seen_real = False
    for idx, c in enumerate(items):
        if isinstance(c, LTChar):
            t = c.get_text()
            font = c.fontname or ""
            pua = bool(t) and "\ue000" <= t[0] <= "\uf8ff"
            if SYMBOL_BULLET_FONT_RE.search(font) or (pua and not ICON_FONT_RE.search(font)):
                chars.append(("• " if not seen_real else " ", c if not seen_real else None))
                continue
            if ICON_FONT_RE.search(font) or t.startswith("(cid:"):
                chars.append((" ", None))
                continue
            # Suppress faux-bold / shadow duplicate glyphs drawn at the same (x0, y0)
            if t.strip() and any(
                t == rc.get_text() and abs(c.x0 - rc.x0) < 1.6 and abs(c.y0 - rc.y0) < 1.6
                for rc in real_seen
            ):
                if chars and chars[-1][1] is None and chars[-1][0] == " ":
                    chars.pop()
                continue
            if t.strip():
                seen_real = True
                real_seen.append(c)
            chars.append((t, c))
        else:  # LTAnno (inferred spaces)
            # Suppress false kerning spaces inside words when adjacent LTChars are tightly spaced
            txt_anno = c.get_text()
            if txt_anno == " " and chars and chars[-1][1] is not None:
                prev_c = chars[-1][1]
                next_c = next((it for it in items[idx + 1:] if isinstance(it, LTChar)), None)
                if next_c is not None and next_c.x0 < prev_c.x1 - 2.0:
                    continue
                if (
                    next_c is not None
                    and prev_c.get_text().isalpha()
                    and next_c.get_text().isalpha()
                    and next_c.get_text().islower()
                    and (next_c.x0 - prev_c.x1) < 0.085 * max(prev_c.size, 8.0)
                ):
                    continue
            chars.append((txt_anno, None))
    real = [c for _, c in chars if c is not None and c.get_text().strip()]
    if not real:
        return None
    text = "".join(t for t, _ in chars).replace("\n", " ")
    # Canva / Figma exports encode word spaces as literal TAB glyphs inside one text run; real column
    # tabs are re-inserted later by _row() from geometry, so in-run tabs are always plain spaces.
    text = re.sub(r"[ \t]*\t[ \t]*", " ", text)
    sizes = [c.size for c in real]
    bold_ratio = sum(1 for c in real if BOLD_FONT_RE.search(c.fontname or "")) / len(real)
    lead_chars = [c for c in real[:14] if c.get_text().isalpha()]
    lead_bold = len(lead_chars) >= 4 and (
        sum(1 for c in lead_chars if BOLD_FONT_RE.search(c.fontname or "")) / len(lead_chars)
    ) > 0.75
    return {
        "text": text,
        "x0": min(c.x0 for c in real), "x1": max(c.x1 for c in real),
        "y0": min(c.y0 for c in real), "y1": max(c.y1 for c in real),
        "size": statistics.median(sizes), "bold": bold_ratio > 0.55 or lead_bold,
    }


def _find_gutter(frags, width) -> Optional[float]:
    """Return x of a vertical gutter splitting the page into two text columns, or None."""
    # Ignore full-width spanning lines (e.g. header/summary across the top) when searching for the column gutter
    body = [f for f in frags if len(f["text"].strip()) > 1 and (f["x1"] - f["x0"]) <= 0.58 * width]
    if len(body) < 10:
        return None
    best = None
    for g in range(int(width * 0.20), int(width * 0.80), 2):
        left = [f for f in body if f["x1"] <= g + 2]
        right = [f for f in body if f["x0"] >= g - 2]
        cross = len(body) - len(left) - len(right)
        if len(left) < 0.15 * len(body) or len(right) < 0.15 * len(body):
            continue
        if cross > max(3, 0.12 * len(body)):
            continue
        # gap width around g
        lmax = max((f["x1"] for f in left), default=0)
        rmin = min((f["x0"] for f in right), default=width)
        score = (cross, -(rmin - lmax))
        if best is None or score < best[0]:
            best = (score, g, left, right)
    if not best:
        return None
    _, g, left, right = best
    # Reject "single column + right-aligned dates": most right fragments share a baseline with a left one
    def has_partner(f, pool):
        return any(abs(f["y1"] - p["y1"]) < max(2.5, 0.35 * f["size"]) for p in pool)
    partnered = sum(1 for f in right if has_partner(f, left))
    right_text = " ".join(f["text"] for f in right)
    datey = len(re.findall(r"\b(19|20)\d{2}\b|present|présent|current|aujourd", right_text, re.I))
    if partnered / max(1, len(right)) > 0.55 and datey >= 0.3 * len(right):
        return None
    if partnered / max(1, len(right)) > 0.8:
        return None
    # Reject "single column + left date/label margin" (Europass / eJobs layouts)
    partnered_l = sum(1 for f in left if has_partner(f, right))
    left_text = " ".join(f["text"] for f in left)
    datey_l = len(re.findall(r"\b(19|20)\d{2}\b|present|présent|current|aujourd|\b\d+\s*(?:years?|months?|ans?|mois)\b", left_text, re.I))
    if g < 0.33 * width and (
        (partnered_l / max(1, len(left)) > 0.40 and datey_l >= 0.25 * len(left))
        or partnered_l / max(1, len(left)) > 0.65
        or datey_l >= 0.45 * len(left)
    ):
        return None
    return float(g)


def _merge_rows(frags, col_of) -> List[Dict]:
    """Merge fragments that share a baseline (within a column) into rows; TAB between distant fragments."""
    rows: List[Dict] = []
    by_col: Dict[int, List[Dict]] = {}
    for f in frags:
        by_col.setdefault(col_of(f), []).append(f)
    for col, fs in by_col.items():
        fs.sort(key=lambda f: (-f["y1"], f["x0"]))
        cur: List[Dict] = []
        for f in fs:
            if cur and abs(cur[0]["y1"] - f["y1"]) <= max(2.2, 0.38 * max(f["size"], cur[0]["size"])):
                cur.append(f)
            else:
                if cur:
                    rows.append(_row(cur, col))
                cur = [f]
        if cur:
            rows.append(_row(cur, col))
    return rows


def _row(fs, col) -> Dict:
    fs = sorted(fs, key=lambda f: (round(f["x0"], 1), -len(f["text"])))
    dedup_fs = [fs[0]]
    for f in fs[1:]:
        prev = dedup_fs[-1]
        t_cur = f["text"].strip()
        t_prev = prev["text"].strip()
        # Suppress faux-bold / partial label overlays starting at the same x0 or overlapping horizontally
        if f["x0"] < prev["x1"] - 2.0 and (t_cur == t_prev or t_prev.startswith(t_cur) or t_cur in t_prev):
            prev["bold"] = prev["bold"] or f["bold"]
            continue
        if f["x0"] < prev["x1"] - 2.0 and t_cur.startswith(t_prev):
            f["bold"] = prev["bold"] or f["bold"]
            dedup_fs[-1] = f
            continue
        dedup_fs.append(f)
    fs = dedup_fs
    parts = [fs[0]["text"].strip()]
    for prev, f in zip(fs, fs[1:]):
        gap = f["x0"] - prev["x1"]
        # Only insert TAB for genuine column/tab stops or right-aligned metadata, never for justified prose gaps
        is_tab = (
            gap > max(20.0, 2.0 * f["size"])
            or (gap > max(9.0, 0.95 * f["size"]) and (f["x0"] > 300 or col == 0 and prev["x0"] < 90 and f["x0"] > 125))
        )
        sep = "\t" if is_tab else " "
        parts.append(sep + f["text"].strip())
    text = "".join(parts)
    biggest = max(fs, key=lambda f: len(f["text"]))
    lead_is_bold = bool(fs[0]["bold"] and len(fs[0]["text"].strip()) >= 3)
    return {
        "text": text, "x0": fs[0]["x0"], "x1": fs[-1]["x1"], "y1": max(f["y1"] for f in fs),
        "y0": min(f["y0"] for f in fs), "size": max(f["size"] for f in fs),
        "bold": lead_is_bold or biggest["bold"] or all(f["bold"] for f in fs), "col": col,
    }


def extract_pdf(pdf_bytes: bytes) -> Document:
    try:
        pages = _fragments_from_pdfminer(pdf_bytes)
    except Exception:
        return _extract_pdf_pypdf2(pdf_bytes)

    raw_chars = sum(len(f["text"].strip()) for _, frags, _ in pages for f in frags)
    used_ocr = False
    if raw_chars < 80:
        ocr_pages = _fragments_from_ocr(pdf_bytes)
        if ocr_pages and sum(len(f["text"].strip()) for _, frags, _ in ocr_pages for f in frags) >= 80:
            pages = ocr_pages
            used_ocr = True
        else:
            try:
                pypdf_doc = _extract_pdf_pypdf2(pdf_bytes)
                if not pypdf_doc.meta.get("scanned"):
                    return pypdf_doc
            except Exception:
                pass

    links = _pdf_links(pdf_bytes)
    lines: List[Line] = []
    offset = 0.0
    columns_seen = 0
    has_photo = False
    total_chars = 0
    for pi, (bbox, frags, images) in enumerate(pages):
        x_min, y_min, width, height = bbox[0], bbox[1], bbox[2] - bbox[0], bbox[3] - bbox[1]
        total_chars += sum(len(f["text"].strip()) for f in frags)
        for (ix0, iy0, ix1, iy1) in images:
            w, h = ix1 - ix0, iy1 - iy0
            if pi == 0 and 45 < w < 220 and 45 < h < 260 and 0.6 < w / max(h, 1) < 1.6 and iy1 > bbox[3] * 0.55:
                has_photo = True
        gutter = None if (pi > 0 and columns_seen == 0) else _find_gutter(frags, width)
        if gutter is not None:
            columns_seen = max(columns_seen, 2)

        def col_of(f, g=gutter, w=width, xm=x_min):
            if g is None:
                return 0
            if f["x1"] <= g + 1:
                return 0
            if f["x0"] >= g - 1:
                return 1
            if (f["x1"] - f["x0"]) > 0.55 * w:
                return -1
            return 0 if (f["x0"] + f["x1"]) / 2 < g else 1

        rows = _merge_rows(frags, col_of)
        page_links = [(r, u) for (p, r, u) in links if p == pi]
        page_row_lines: List[Tuple[Dict, Line]] = []
        for r in rows:
            text = normalize_text(r["text"])
            if not text:
                continue
            if len(pages) > 1 and r["y0"] < bbox[1] + 45 and re.search(
                r"\b(?:page\s*\d+|\d+\s*/\s*\d+|ejobs\.ro|curriculum\s+vitae)\b",
                text,
                re.I,
            ):
                continue
            ln = Line(
                text=text, size=round(r["size"], 1), bold=r["bold"], x0=r["x0"], x1=r["x1"],
                top=offset + (bbox[3] - r["y1"]), page=pi, col=r["col"], caps=_is_caps(text),
            )
            page_row_lines.append((r, ln))
            lines.append(ln)

        # Attach each PDF link annotation to the single best matching row on this page
        for (rx0, ry0, rx1, ry1), uri in page_links:
            clean_u = sanitize_url(uri)
            if not clean_u or "ejobs.ro" in clean_u.lower():
                continue
            annot_ymid = 0.5 * (ry0 + ry1)
            best_ln = None
            best_dy = 999.0
            for r, ln in page_row_lines:
                if rx0 <= r["x1"] + 4 and rx1 >= r["x0"] - 4:
                    row_ymid = 0.5 * (r["y0"] + r["y1"])
                    dy = abs(row_ymid - annot_ymid)
                    if dy <= max(6.0, 0.75 * r["size"]) and dy < best_dy:
                        best_dy = dy
                        best_ln = ln
            if best_ln is not None and clean_u not in best_ln.links:
                best_ln.links.append(clean_u)
        offset += height

    # Reading order: per page, spanning rows above the columns, then left column, then right column
    lines = _postprocess(_order_columns(lines))
    sizes = [l.size for l in lines for _ in range(max(1, len(l.text) // 20))]
    body = statistics.median(sizes) if sizes else 10.0
    doc = Document(
        lines=lines, body_size=body,
        links=[u for u in (sanitize_url(raw_u) for _, _, raw_u in links) if u and "ejobs.ro" not in u.lower()],
        meta={
            "source": "pdf", "pages": len(pages), "layout": "two-column" if columns_seen >= 2 else "single-column",
            "has_photo": has_photo, "scanned": total_chars < 80, "ocr_recovered": used_ocr,
            "extractor": "vision-ocr" if used_ocr else "pdfminer",
        },
    )
    return doc


_BLOCKED_DOMAINS = (
    "job-search-two-beryl.vercel.app",
    # webmail inbox links that Canva/Word users accidentally attach to their e-mail text
    "mail.google.com",
    "outlook.live.com",
    "outlook.office.com",
    "mail.yahoo.com",
)


def sanitize_url(uri: str) -> Optional[str]:
    """Repair common template bugs and filter blocked URLs."""
    u = (uri or "").strip()
    if not u or u.lower().startswith(("mailto:", "tel:")):
        return u or None
    m = re.search(r"(https?://)(?:[^/]+/(?:in/)?)?(https?://.+)$", u)
    if m:
        u = m.group(2)
    u = re.sub(r"(linkedin\.com/in/)(?:https?://)?(?:www\.)?linkedin\.com/in/", r"\1", u, flags=re.I)
    u = re.sub(r"(github\.com/)(?:https?://)?(?:www\.)?github\.com/", r"\1", u, flags=re.I)
    if " " in u.strip():
        return None
    u_low = u.lower()
    if any(b in u_low for b in _BLOCKED_DOMAINS):
        return None
    return u


_PAGE_NO_RE = re.compile(r"^(?:page\s*)?\d{1,2}(?:\s*/\s*\d{1,2})?$", re.I)


def _postprocess(lines: List[Line]) -> List[Line]:
    out: List[Line] = []
    pending_bullet: Optional[Line] = None
    for l in lines:
        t = l.text.strip()
        if _PAGE_NO_RE.match(t) or t in {"/", "|", "-"}:
            continue
        if all(c in BULLET_GLYPHS + "-*\t " for c in t):
            pending_bullet = l
            continue
        if pending_bullet is not None and pending_bullet.col == l.col and pending_bullet.page == l.page:
            l.text = "• " + t
        pending_bullet = None
        l.links = [u for u in (sanitize_url(x) for x in l.links) if u]
        # Clean inline blocked URLs or orphaned "Live demo · " when live demo URL was stripped
        for bd in _BLOCKED_DOMAINS:
            if bd in l.text.lower():
                l.text = re.sub(rf"(?:https?://)?{re.escape(bd)}/?", "", l.text, flags=re.I).strip(" •·|")
        out.append(l)
    return out


def _order_columns(lines: List[Line]) -> List[Line]:
    """Per page, rows spanning both columns cut the page into horizontal bands; each band is
    read left column first, then right column. Keeps a full-width summary between header and body."""
    out: List[Line] = []
    for p in sorted({l.page for l in lines}):
        pl = [l for l in lines if l.page == p]
        if not any(l.col == 1 for l in pl):
            out.extend(sorted(pl, key=lambda l: (l.top, l.x0)))
            continue
        spans = sorted([l for l in pl if l.col == -1], key=lambda l: l.top)
        cuts = [s.top for s in spans] + [float("inf")]
        prev = float("-inf")
        span_iter = iter(spans)
        for cut in cuts:
            band = [l for l in pl if l.col in (0, 1) and prev <= l.top < cut]
            out.extend(sorted([l for l in band if l.col == 0], key=lambda l: (l.top, l.x0)))
            out.extend(sorted([l for l in band if l.col == 1], key=lambda l: (l.top, l.x0)))
            if cut != float("inf"):
                out.append(next(span_iter))
            prev = cut
    return out


def _extract_pdf_pypdf2(pdf_bytes: bytes) -> Document:
    import PyPDF2
    reader = PyPDF2.PdfReader(io.BytesIO(pdf_bytes))
    raw = "\n".join(p.extract_text() or "" for p in reader.pages)
    lines = [Line(text=normalize_text(t)) for t in raw.splitlines() if normalize_text(t)]
    for l in lines:
        l.caps = _is_caps(l.text)
    return Document(
        lines=lines, links=[u for _, _, u in _pdf_links(pdf_bytes)],
        meta={"source": "pdf", "pages": len(reader.pages), "layout": "unknown", "has_photo": False,
              "scanned": len(raw.strip()) < 80, "extractor": "pypdf2"},
    )


# ─── DOCX / TXT ───────────────────────────────────────────────────────────────

def extract_docx(data: bytes) -> Document:
    import docx  # python-docx
    d = docx.Document(io.BytesIO(data))
    lines: List[Line] = []
    links: List[str] = []
    try:
        for rel in d.part.rels.values():
            if "hyperlink" in rel.reltype and rel.target_ref:
                links.append(rel.target_ref)
    except Exception:
        pass

    def add_par(p):
        text = normalize_text(p.text)
        if not text:
            return
        runs = [r for r in p.runs if r.text.strip()]
        sizes = [r.font.size.pt for r in runs if r.font is not None and r.font.size]
        bold = bool(runs) and all((r.bold or (p.style is not None and p.style.font is not None and p.style.font.bold)) for r in runs)
        style = (p.style.name if p.style is not None else "") or ""
        size = max(sizes) if sizes else (14.0 if style.lower().startswith(("heading", "title")) else 10.5)
        if style.lower().startswith("list") and not text.startswith(tuple(BULLET_GLYPHS + "-*")):
            text = "• " + text
        lines.append(Line(text=text, size=size, bold=bold or style.lower().startswith(("heading", "title")), caps=_is_caps(text)))

    for block in d.element.body.iterchildren():
        tag = block.tag.split("}")[-1]
        if tag == "p":
            add_par(docx.text.paragraph.Paragraph(block, d))
        elif tag == "tbl":
            tbl = docx.table.Table(block, d)
            for row in tbl.rows:
                seen = set()
                cells_text = []
                for cell in row.cells:
                    if id(cell._tc) in seen:
                        continue
                    seen.add(id(cell._tc))
                    for p in cell.paragraphs:
                        add_par(p)
    body = statistics.median([l.size for l in lines]) if lines else 10.5
    return Document(lines=lines, body_size=body, links=links,
                    meta={"source": "docx", "pages": 1, "layout": "single-column", "has_photo": False, "scanned": False, "extractor": "python-docx"})


def extract_txt(text: str) -> Document:
    lines = [Line(text=normalize_text(t)) for t in (text or "").splitlines() if normalize_text(t)]
    for l in lines:
        l.caps = _is_caps(l.text)
    return Document(lines=lines, meta={"source": "text", "pages": 1, "layout": "single-column", "has_photo": False, "scanned": False, "extractor": "text"})


def extract_document(data: Optional[bytes] = None, filename: str = "", text: str = "") -> Document:
    """Entry point: bytes + filename (pdf/docx/txt) or raw text."""
    name = (filename or "").lower()
    if isinstance(data, str):
        data = data.encode("utf-8")
    if data:
        head = data.lstrip()[:8]
        if head.startswith(b"%PDF-"):
            doc = extract_pdf(data)
            if not doc.meta.get("scanned") or not text:
                return doc
        elif data[:4] == b"PK\x03\x04" or name.endswith(".docx"):
            try:
                return extract_docx(data)
            except Exception:
                pass
        elif name.endswith(".txt") or not head.startswith(b"%PDF-"):
            decoded = data.decode("utf-8", errors="ignore")
            if decoded.strip():
                return extract_txt(decoded)
    return extract_txt(text)
